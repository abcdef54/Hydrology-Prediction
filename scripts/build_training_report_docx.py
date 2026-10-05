"""Build the editable Vietnamese report from the verified evaluation exports."""

import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
EXPORTS = ROOT / "reports/model_training"
OUTPUT = ROOT / "Báo Cáo Train models.docx"
MODELS = {
    "lgbm": "LightGBM", "xgb": "XGBoost", "lstm": "Vanilla LSTM",
    "residual_lstm": "Residual LSTM", "lstm_adapter": "MEF-LSTM",
    "persistence": "Persistence",
}
HORIZONS = [60, 180, 360, 720, 1440]


def read_csv(name):
    with (EXPORTS / name).open(encoding="utf-8", newline="") as source:
        return list(csv.DictReader(source))


def number(value, digits=2):
    return f"{float(value):.{digits}f}".replace(".", ",")


def paragraph(document, text, style=None):
    return document.add_paragraph(text, style)


def table(document, headers, rows, best_cells=()):
    result = document.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.autofit = False
    widths = [3.3] + [(17 - 3.3) / (len(headers) - 1)] * (len(headers) - 1)
    for column, width in zip(result.columns, widths):
        column.width = Cm(width)
    for cell, label in zip(result.rows[0].cells, headers):
        cell.text = label
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E8EDF2")
        cell._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement("w:tblHeader")
    result.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, values in enumerate(rows):
        cells = result.add_row().cells
        for column_index, value in enumerate(values):
            cells[column_index].text = str(value)
            if (row_index, column_index) in best_cells:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E2EFDA")
                cells[column_index]._tc.get_or_add_tcPr().append(shade)
    for row_index, row in enumerate(result.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for column_index, cell in enumerate(row.cells):
            for cell_paragraph in cell.paragraphs:
                cell_paragraph.paragraph_format.space_after = Pt(3)
                cell_paragraph.paragraph_format.space_before = Pt(3)
                cell_paragraph.paragraph_format.line_spacing = 1
                if column_index:
                    cell_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in cell_paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = row_index == 0 or (row_index - 1, column_index) in best_cells
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return result


def metric_table(document, records, metrics, labels, digits):
    by_model = {record["model"]: record for record in records}
    minima = []
    for metric in metrics:
        values = [float(by_model[model][metric]) for model in MODELS]
        minima.append(max(values) if metric in {"nse", "kge"} else min(values))
    rows, best_cells = [], []
    for row_index, (model, name) in enumerate(MODELS.items()):
        record = by_model[model]
        rows.append([name] + [number(record[key], precision) for key, precision in zip(metrics, digits)])
        for column_index, metric in enumerate(metrics, 1):
            if float(record[metric]) == minima[column_index - 1]:
                best_cells.append((row_index, column_index))
    return table(document, ["Mô hình"] + labels, rows, best_cells)


def horizon_table(document, records, metric="rmse"):
    by_key = {(record["model"], record["target"]): record for record in records}
    rows, best_cells = [], []
    for row_index, (model, name) in enumerate(MODELS.items()):
        values = []
        for column_index, horizon in enumerate(HORIZONS, 1):
            target = f"target_water_level_plus_{horizon}m"
            value = float(by_key[model, target][metric])
            comparison = [float(by_key[other, target][metric]) for other in MODELS]
            best = max(comparison) if metric == "nse" else min(comparison)
            if value == best:
                best_cells.append((row_index, column_index))
            values.append(number(value, 3 if metric == "nse" else 2))
        rows.append([name] + values)
    return table(document, ["Mô hình"] + [f"{h:,} phút".replace(",", ".") for h in HORIZONS], rows, best_cells)


def build_report():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in [("Title", 20), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Trang ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    title = paragraph(document, "BÁO CÁO HUẤN LUYỆN\nDự báo mực nước tại trạm Dã Viên", "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author = paragraph(document, "Nguyen Minh Huy • Tháng 10 năm 2026")
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("I. Mục tiêu huấn luyện", 1)
    document.add_heading("1. Bài toán cần giải quyết", 2)
    paragraph(document, "Mục tiêu là dự báo mực nước sông Hương tại trạm Dã Viên sau 60, 180, 360, 720 và 1.440 phút dựa trên các số liệu đã quan sát được. Báo cáo so sánh năm phương pháp: LightGBM, XGBoost, Vanilla LSTM, Residual LSTM và MEF-LSTM, cùng một mô hình cơ sở giữ nguyên mực nước.")
    paragraph(document, "Mỗi phương pháp được thử với ba nhóm đầu vào: mực nước, mực nước + mưa và tất cả biến. Dữ liệu được lấy theo giờ. Thiết lập horizon=1h có nghĩa là mỗi bước dữ liệu cách nhau một giờ; mô hình vẫn dự báo cho cả năm thời hạn nêu trên.")
    document.add_heading("II. Dữ liệu huấn luyện", 1)
    document.add_heading("1. Tiền xử lý dữ liệu", 2)
    for text in [
        "Dữ liệu từ các nguồn được ghép theo thời gian UTC. Mực nước được tính trung bình theo giờ, còn lượng mưa được cộng theo giờ. Các biến đầu vào bổ sung gồm mực nước của 1 và 3 giờ trước, cùng tổng lượng mưa trong 1, 3, 6 và 24 giờ gần nhất.",
        "Các cờ đánh dấu dữ liệu thiếu được giữ lại. Với hồ chứa, giá trị quá khứ chỉ được dùng để điền vào chỗ thiếu trong tối đa 60 phút. Cách điền này không được áp dụng cho mực nước, mưa và gió. Dữ liệu sau đó được sắp xếp theo thời gian; các thời điểm trùng và các dòng thiếu một trong năm giá trị cần dự báo được loại bỏ.",
        "Đối với LSTM, dữ liệu đầu vào và đầu ra được chuẩn hóa bằng trung bình và độ lệch chuẩn của tập Train. Các giá trị thống kê này cũng được dùng để chuẩn hóa tập Val và Test. Những ô đầu vào còn thiếu được thay bằng 0 sau khi chuẩn hóa, đồng thời giữ lại cờ đánh dấu dữ liệu thiếu. Số 0 ở đây tương ứng với mức trung bình của tập Train, không có nghĩa là mực nước hoặc lượng mưa bằng 0. LightGBM và XGBoost có thể xử lý trực tiếp giá trị thiếu nên không cần thay bằng 0.",
    ]:
        paragraph(document, text)
    paragraph(document, "Bộ dữ liệu theo giờ có 7.667 dòng. Nhóm mực nước gồm 4 biến: mực nước hiện tại, mực nước trễ 60 phút, trễ 180 phút và cờ thiếu mực nước. Nhóm mực nước + mưa có 88 biến; nhóm tất cả có 176 biến, gồm thêm dữ liệu mưa của 14 trạm, gió tại Thuận An/Tư Hiền và 4 hồ chứa.")
    document.add_heading("2. Cách chia Train/Val/Test", 2)
    paragraph(document, "Dữ liệu được chia theo thứ tự thời gian với tỷ lệ ban đầu 70% / 15% / 15%. Các dòng ở cuối tập Train và Val được loại bỏ nếu giá trị cần dự báo sau 24 giờ nằm trong tập tiếp theo, nhằm tránh dùng thông tin tương lai. Tập Val được dùng để chọn cấu hình mô hình tốt nhất, xác định thời điểm dừng sớm và chọn phiên bản mô hình được lưu lại. Tập Test được dùng để đánh giá kết quả cuối cùng. Các mốc thời gian trong bảng dưới đây đều theo UTC.")
    table(document, ["Tập", "Số dòng", "Khoảng thời gian (UTC)"], [
        ["Train", "5.326", "14/11/2025 18:00 – 24/06/2026 15:00"],
        ["Val", "1.122", "25/06/2026 16:00 – 11/08/2026 09:00"],
        ["Test", "1.147", "12/08/2026 10:00 – 29/09/2026 04:00"],
    ])
    document.add_heading("3. Số lượng dữ liệu huấn luyện", 2)
    paragraph(document, "LightGBM và XGBoost sử dụng toàn bộ số dòng trong bảng trên. LSTM cần một đoạn dữ liệu quá khứ để tạo mỗi mẫu, vì vậy số mẫu ít hơn. Mỗi mẫu được dùng để dự báo năm giá trị mực nước tương lai.")
    table(document, ["Chuỗi lịch sử", "Train", "Val", "Test"], [["72 giờ", "5.255", "1.051", "1.076"], ["144 giờ", "5.183", "979", "1.004"]])
    paragraph(document, "Thông tin ngày tháng chỉ được dùng để chia tập và đánh giá kết quả, chưa được đưa vào mô hình dưới dạng các biến thể hiện mùa hoặc thời điểm trong năm.")

    document.add_page_break()
    document.add_heading("III. Các mô hình huấn luyện", 1)
    paragraph(document, "Thông số được lấy từ các lần chạy đã hoàn tất trong MLflow, experiment 1h (ID 3), với seed 42. LightGBM và XGBoost được huấn luyện riêng cho từng thời hạn, còn LSTM dự báo cả năm thời hạn cùng lúc. LR là tốc độ cập nhật trọng số. Patience là số epoch liên tiếp mà kết quả trên tập Val không cải thiện trước khi quá trình huấn luyện dừng lại. Các LSTM giảm LR còn 90% sau mỗi 10 epoch và giới hạn chuẩn gradient ở 1.")
    document.add_heading("1. LightGBM", 2)
    paragraph(document, "LightGBM cộng dần nhiều cây quyết định, mỗi cây mới giúp giảm sai số còn lại. Mô hình sử dụng các biến ở từng dòng dữ liệu, bao gồm các giá trị của những thời điểm trước đã được tạo ở bước tiền xử lý.")
    paragraph(document, "Cấu hình: tối đa 2.000 cây; LR 0,03; tối đa 15 lá/cây; độ sâu 6; tỷ lệ lấy mẫu dòng và cột 0,8; lấy mẫu dòng mỗi vòng; L1 = 0,5; L2 = 1,5. Quá trình huấn luyện dừng sớm nếu kết quả trên tập Val không cải thiện sau 100 vòng liên tiếp.")
    document.add_heading("2. XGBoost", 2)
    paragraph(document, "XGBoost cũng cộng dần các cây để giảm sai số, đồng thời dùng các hệ số phạt để hạn chế mô hình quá phức tạp.")
    paragraph(document, "Cấu hình: tối đa 2.000 cây; LR 0,03; độ sâu 6; min child weight = 5; tỷ lệ lấy mẫu dòng và cột 0,75; L1 = 1,0; L2 = 1,5. Quá trình huấn luyện dừng sớm nếu kết quả trên tập Val không cải thiện sau 100 vòng liên tiếp.")
    document.add_heading("3. Vanilla LSTM", 2)
    paragraph(document, "LSTM đọc chuỗi quan sát trong quá khứ, lấy trạng thái ở bước cuối và dự báo trực tiếp năm mực nước tương lai.")
    paragraph(document, "Cấu hình: 3 lớp, hidden size 32, chuỗi 144 giờ, batch size 128; tối đa 150 epoch, patience 50. Bộ tối ưu AdamW, LR 0,01, weight decay 0,001; dropout giữa các lớp 0,5, input dropout 0,1 và head dropout 0,2. Gradient accumulation = 1; huấn luyện trên CUDA.")
    document.add_heading("4. Residual LSTM", 2)
    paragraph(document, "Đây là kiến trúc LSTM tùy chỉnh: mô hình học mức tăng hoặc giảm so với mực nước hiện tại. Dự báo cuối cùng bằng mực nước vừa quan sát cộng với phần thay đổi mà LSTM dự đoán:")
    equation = paragraph(document, "Mực nước dự báo = Mực nước hiện tại + Phần thay đổi dự báo")
    equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(document, "Nếu phần thay đổi bằng 0, mô hình giữ nguyên mực nước như baseline. Trong mã nguồn, hai thành phần được đưa về cùng thang chuẩn hóa của từng đích trước khi cộng.")
    paragraph(document, "Loss huấn luyện gồm MSE cộng 0,1 lần trung bình bình phương phần thay đổi trên thang chuẩn hóa. Khoản phạt này giúp tránh hiệu chỉnh quá lớn, nhưng cũng có thể làm dự báo thấp khi nước tăng mạnh. Khi đánh giá trên tập Val và Test, chỉ sai số dự báo được tính; khoản phạt này không được cộng vào kết quả.")
    paragraph(document, "Thông số chung: 3 lớp, hidden size 16, batch size 128; dropout giữa lớp 0,5, input dropout 0,1, head dropout 0,2; correction penalty 0,1; AdamW, gradient accumulation = 1 và CUDA.")
    table(document, ["Nhóm đầu vào", "Chuỗi (giờ)", "LR", "Weight decay", "Epoch tối đa", "Patience"], [["Mực nước", "144", "0,009", "0,005", "100", "30"], ["Mực nước + mưa / Tất cả", "72", "0,01", "0,001", "150", "50"]])
    paragraph(document, "Với nhóm đầu vào mực nước, mô hình được chọn có MSE trên tập Val thấp nhất trong ba lần chạy đang lưu, đạt 0,03084. Mô hình dùng chuỗi 72 giờ đạt 0,03207. Do các nhóm đầu vào sử dụng cấu hình khác nhau, kết quả hiện tại chưa cho biết riêng việc thêm biến ảnh hưởng đến chất lượng dự báo như thế nào.")
    document.add_heading("5. MEF-LSTM (Mean Embedding Forecast LSTM)", 2)
    paragraph(document, "Mô hình dùng phần LSTM nền đã huấn luyện trước trong model_epoch085.pt, hidden size 512. Adapter chuyển các biến địa phương sang đầu vào của mô hình nền; lớp đầu ra mới dự báo năm đích. Checkpoint gốc học lưu lượng, còn bài toán này dự báo mực nước.")
    paragraph(document, "Đây là phiên bản MEF điều chỉnh cho Dã Viên. Do không có dự báo thời tiết tương lai, embedding tương ứng được đặt bằng 0; embedding và lớp đầu ra gốc không được dùng trong đường dự báo hiện tại.")
    paragraph(document, "Cấu hình chung: chuỗi 144 giờ, batch size 128; AdamW, weight decay 0,001; input dropout 0,05, head dropout 0,3; patience 30, gradient accumulation = 1 và CUDA.")
    paragraph(document, "Giai đoạn 1: giữ cố định mô hình nền, chỉ học adapter và lớp đầu ra; LR = 3 × 10⁻⁴, tối đa 100 epoch. Giai đoạn 2: nạp bản tốt nhất của giai đoạn 1, cho phép cập nhật cả hai LSTM nền; LR nền = 5 × 10⁻⁴, LR adapter/đầu ra = 10⁻³, tối đa 50 epoch.")
    paragraph(document, "Phiên bản có sai số thấp nhất trên tập Val trong cả hai giai đoạn được chọn để đánh giá trên tập Test. Với nhóm mực nước và mực nước + mưa, phiên bản được chọn thuộc giai đoạn 2. Với nhóm tất cả biến, phiên bản tốt nhất vẫn thuộc giai đoạn 1.")
    document.add_heading("6. Mô hình cơ sở: giữ nguyên mực nước (Persistence)", 2)
    paragraph(document, "Persistence lấy mực nước vừa quan sát làm dự báo cho cả năm thời hạn. Ví dụ, nếu hiện tại là H, dự báo sau 1, 3, 6, 12 và 24 giờ đều bằng H. Mô hình không học tham số và không dùng mưa, gió hay hồ chứa.")
    paragraph(document, "Baseline này cho biết mô hình huấn luyện có tốt hơn một dự báo đơn giản hay không. Hạn chế của mô hình là không dự báo trước được mức tăng hoặc giảm của mực nước, nên đỉnh lũ thường được dự báo muộn.")

    details = read_csv("test_metrics_common_rows.csv")
    summaries = read_csv("test_metrics_summary.csv")
    document.add_page_break()
    document.add_heading("IV. Kết quả huấn luyện", 1)
    document.add_heading("1. Cách đọc kết quả", 2)
    paragraph(document, "Các mô hình đã lưu được đánh giá lại trên cùng 1.004 thời điểm của tập Test, từ 18/08/2026 09:00 đến 29/09/2026 04:00 UTC. Persistence cũng được đánh giá trên các thời điểm này. Trước khi so sánh trên các dòng dữ liệu chung, sai số của từng mô hình đã được tính lại trên tập đánh giá gốc và kiểm tra khớp với kết quả trong MLflow.")
    paragraph(document, "RMSE, MAE và MSE càng thấp càng tốt. MSE là trung bình bình phương sai số, RMSE là căn của MSE, còn MAE là trung bình độ lớn sai số. NSE càng gần 1 càng tốt. NSE bằng 0 có nghĩa là mô hình tương đương với việc luôn dự báo bằng mực nước trung bình của tập đánh giá; NSE âm cho thấy mô hình còn kém hơn cách dự báo này. KGE càng gần 1 càng tốt, phản ánh độ khớp diễn biến, độ dao động và mức trung bình.")
    paragraph(document, "Bảng RMSE theo thời hạn cho biết mô hình tốt nhất tại từng mốc. Các bảng tổng hợp lấy trung bình MSE/MAE với trọng số bằng nhau trên năm thời hạn; RMSE tổng hợp là căn của MSE đó, không phải trung bình các RMSE. NSE và KGE tổng hợp được tính bằng trung bình các chỉ số tương ứng của năm thời hạn.")
    paragraph(document, "Sai số được giữ ở đơn vị gốc của dữ liệu; MSE có đơn vị bình phương. Do đơn vị vật lý và mốc đo chưa được xác nhận, báo cáo chưa ghi sai số theo mét hoặc centimét. Giá trị test_loss được lưu trong MLflow cũng không thể dùng để so sánh trực tiếp mô hình cây với LSTM, vì mô hình cây tính sai số trên thang đo gốc, còn LSTM tính trên dữ liệu đã chuẩn hóa.")
    paragraph(document, "Các ô được in đậm và tô nền xanh thể hiện kết quả tốt nhất trong từng cột, bao gồm cả Persistence. Thời hạn được tính từ lúc phát dự báo; 60 / 180 / 360 / 720 / 1.440 phút tương ứng 1 / 3 / 6 / 12 / 24 giờ.")
    feature_notes = [
        ("water_level", "2. Mực nước (Water Level)", "Residual LSTM tốt nhất ở 60, 180, 360 và 720 phút. Ở 1.440 phút, Persistence có RMSE thấp nhất (102,26); Residual LSTM đứng đầu trong năm mô hình huấn luyện (110,02)."),
        ("water_level+rain", "3. Mực nước + mưa (Water Level + Rain)", "XGBoost tốt nhất ở 60 phút. Residual LSTM có RMSE thấp nhất ở thời hạn 180 phút, nhưng chỉ thấp hơn Persistence khoảng 0,02 đơn vị. Từ 360 đến 1.440 phút, Persistence tốt nhất; thêm mưa chưa cải thiện sai số tổng hợp trong các cấu hình này."),
        ("all", "4. Tất cả biến", "XGBoost tốt nhất ở 60 và 180 phút; LightGBM tốt nhất ở 360 phút. Residual LSTM có RMSE thấp nhất ở thời hạn 720 phút, nhưng chỉ thấp hơn Persistence khoảng 0,02 đơn vị. Ở 1.440 phút, Persistence tiếp tục tốt nhất. Các chênh lệch rất nhỏ chưa đủ để kết luận mô hình vượt trội."),
    ]
    for feature_set, heading, note in feature_notes:
        document.add_page_break()
        document.add_heading(heading, 2)
        paragraph(document, "RMSE tại từng thời hạn dự báo (càng thấp càng tốt):")
        horizon_table(document, [r for r in details if r["feature_set"] == feature_set])
        paragraph(document, note)
        paragraph(document, "Kết quả tổng hợp trên năm thời hạn:")
        metric_table(document, [r for r in summaries if r["feature_set"] == feature_set], ["rmse", "mae", "mse", "nse", "kge"], ["RMSE ↓", "MAE ↓", "MSE ↓", "NSE ↑", "KGE ↑"], [2, 2, 2, 3, 3])
        if feature_set == "water_level":
            paragraph(document, "NSE tại từng thời hạn dự báo của nhóm mực nước:")
            horizon_table(document, [r for r in details if r["feature_set"] == feature_set], "nse")

    document.add_page_break()
    document.add_heading("5. Phân loại vượt mức báo động lũ", 2)
    paragraph(document, "Các mô hình được huấn luyện để dự báo mực nước. Cảnh báo lũ được xác định bằng cách so sánh mực nước dự báo tại từng thời hạn với các mức báo động sau:")
    table(document, ["Cấp báo động", "BĐ1", "BĐ2", "BĐ3"], [["Ngưỡng (m)", "2,0", "2,5", "3,5"]])
    paragraph(document, "Báo cáo dùng các ngưỡng do người thực hiện cung cấp cho Dã Viên theo tham chiếu Kim Long. Phép quy chiếu mốc đo giữa hai trạm chưa được xác nhận.")
    paragraph(document, "Với một ngưỡng T, quan sát thực tế được gán nhãn vượt ngưỡng khi mực nước đo ≥ T; dự báo được gán nhãn tương tự khi mực nước dự báo ≥ T. Kết quả được đánh giá riêng cho từng cấp báo động. Nếu chỉ cần phân biệt có hoặc không có cảnh báo từ BĐ1 trở lên, ngưỡng được dùng là T = 2,0 m. Mực nước vượt mức báo động tại trạm không có nghĩa là mọi khu dân cư đều bị ngập.")
    paragraph(document, "Precision cho biết tỷ lệ cảnh báo đúng trong các cảnh báo phát ra. Recall cho biết tỷ lệ lần vượt ngưỡng được phát hiện. F1 cân bằng hai tỷ lệ này; cả ba càng cao càng tốt.")
    paragraph(document, "Các giá trị flood_f1 đang lưu trong MLflow dùng phân vị 90% của mực nước Test, chưa dùng BĐ1–BĐ3. Vì chưa xác nhận cách đổi đơn vị gốc sang mét và mốc đo, báo cáo chưa đưa ra F1, Precision hay Recall theo các mức báo động. Các chỉ số này cần được tính lại từ những dự báo đã lưu sau khi đơn vị và mốc đo được xác nhận.")
    document.add_heading("V. Nhận xét huấn luyện", 1)
    for text in [
        "Residual LSTM với đầu vào mực nước có kết quả tổng hợp tốt nhất: RMSE 77,91, NSE 0,726 và KGE 0,786. MSE thấp hơn Persistence khoảng 39,7%. Vanilla LSTM với mực nước đứng tiếp theo về MSE/RMSE.",
        "Kết quả thay đổi theo thời hạn. Với nhóm mực nước, Residual LSTM tốt nhất từ 60 đến 720 phút, nhưng Persistence tốt hơn ở 1.440 phút. Vì vậy, việc lựa chọn mô hình cần dựa vào thời hạn dự báo cần sử dụng, đồng thời xem xét kết quả tổng hợp.",
        "Persistence là một mốc so sánh khá mạnh. Các cấu hình dùng thêm mưa hoặc tất cả biến chưa giảm MSE tổng hợp so với baseline; một số mô hình vẫn tốt hơn ở các thời hạn riêng lẻ.",
        "Trong các lần chạy này, cả năm phương pháp đều có sai số tổng hợp lớn hơn khi thêm mưa so với chỉ dùng mực nước. Điều đó chưa chứng minh mưa không hữu ích: cấu hình giữa các nhóm còn khác nhau, dữ liệu gió và hồ chứa cũng thiếu nhiều.",
        "MEF-LSTM có điểm xuất phát khác bài toán hiện tại: mô hình gốc học lưu lượng với đầu vào khác. Cách huấn luyện hai giai đoạn giúp giảm sai số trên tập Val, nhưng MSE trên tập Test vẫn cao hơn các LSTM được huấn luyện cho bài toán này và Persistence. Kết quả hiện tại chưa đủ để xác định nguyên nhân cụ thể.",
        "Các mô hình sử dụng độ dài chuỗi lịch sử và cấu hình khác nhau; phần lớn các thí nghiệm cũng chỉ dùng seed 42. Để kết luận chắc chắn hơn, cần thử thêm seed và các giai đoạn thời gian khác. Tập Val nên tiếp tục được dùng để chọn cấu hình, còn tập Test được giữ lại để đánh giá cuối cùng.",
    ]:
        paragraph(document, text)
    paragraph(document, "Nguồn số liệu: các mô hình đã lưu của experiment 1h và bộ đánh giá chung xuất ngày 05/10/2026. Trong mỗi phương pháp và nhóm đầu vào, phiên bản LSTM được chọn dựa trên MSE thấp nhất của tập Val. Số liệu chi tiết và mã lần chạy được lưu trong reports/model_training.")
    document.core_properties.title = "Báo cáo huấn luyện – Dự báo mực nước tại trạm Dã Viên"
    document.core_properties.author = "Nguyen Minh Huy"
    document.core_properties.subject = "So sánh mô hình và RMSE theo thời hạn trên cùng tập Test"
    document.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    build_report()
